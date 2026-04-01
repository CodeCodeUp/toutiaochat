#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt
except ImportError as exc:
    raise SystemExit("python-docx is required. Install it with: pip install python-docx") from exc


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Export a normalized Toutiao content payload to DOCX."
    )
    parser.add_argument("input", help="Path to the normalized content JSON file.")
    parser.add_argument(
        "--output",
        help="Output DOCX path. Defaults to <input-stem>.docx beside the input file.",
    )
    parser.add_argument(
        "--body-font",
        default="SimSun",
        help="Font name for body content. Default: SimSun",
    )
    parser.add_argument(
        "--title-font",
        default="SimHei",
        help="Font name for headings and title. Default: SimHei",
    )
    return parser.parse_args()


def load_payload(path: Path) -> dict:
    try:
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
    except FileNotFoundError as exc:
        raise SystemExit(f"Input file not found: {path}") from exc
    except json.JSONDecodeError as exc:
        raise SystemExit(f"Invalid JSON in {path}: {exc}") from exc

    if not isinstance(payload, dict):
        raise SystemExit("Expected a JSON object payload.")
    return payload


def set_run_font(run, font_name: str, size: int | None = None, bold: bool | None = None) -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def split_blocks(content: str) -> list[str]:
    blocks = re.split(r"\n\s*\n", content.strip())
    return [block.strip() for block in blocks if block.strip()]


def strip_inline_markdown(text: str) -> str:
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"\*([^*]+)\*", r"\1", text)
    return text.strip()


def extract_images(payload: dict) -> list[dict]:
    raw_images = payload.get("images")
    if isinstance(raw_images, list) and raw_images:
        return raw_images

    derived: list[dict] = []
    seen: set[str] = set()

    image_prompts = payload.get("image_prompts", [])
    if isinstance(image_prompts, list):
        for prompt in image_prompts:
            if not isinstance(prompt, dict):
                continue
            position = str(prompt.get("position", "end")).strip() or "end"
            caption = str(prompt.get("caption") or prompt.get("description") or "").strip()
            generated_images = prompt.get("generated_images", [])
            if not isinstance(generated_images, list):
                continue
            for generated in generated_images:
                if not isinstance(generated, dict):
                    continue
                file_path = str(generated.get("file_path") or generated.get("path") or "").strip()
                if not file_path or file_path in seen:
                    continue
                seen.add(file_path)
                derived.append(
                    {
                        "path": file_path,
                        "position": position,
                        "caption": caption,
                        "prompt": caption,
                    }
                )

    generated_image_paths = payload.get("generated_image_paths", [])
    if isinstance(generated_image_paths, list):
        for index, item in enumerate(generated_image_paths):
            file_path = str(item).strip()
            if not file_path or file_path in seen:
                continue
            seen.add(file_path)
            derived.append(
                {
                    "path": file_path,
                    "position": "cover" if index == 0 else "end",
                    "caption": "",
                    "prompt": "",
                }
            )

    return derived


def normalize_images(raw_images: object, paragraph_count: int) -> dict:
    grouped = {"cover": [], "after_paragraph": {}, "end": []}
    if not isinstance(raw_images, list):
        return grouped

    for item in raw_images:
        if not isinstance(item, dict):
            continue
        path = Path(str(item.get("path", "")).strip())
        if not path.exists():
            continue

        image_record = {
            "path": path,
            "caption": str(item.get("caption") or item.get("prompt") or "").strip(),
        }

        position = str(item.get("position", "end")).strip()
        if position == "cover":
            grouped["cover"].append(image_record)
            continue
        if position.startswith("after_paragraph:"):
            suffix = position.split(":", 1)[1].strip()
            if suffix.isdigit():
                number = int(suffix)
                if 1 <= number <= max(paragraph_count, 1):
                    grouped["after_paragraph"].setdefault(number, []).append(image_record)
                    continue
        grouped["end"].append(image_record)

    return grouped


def add_image(document: Document, image_path: Path, caption: str, width_inches: float = 5.5) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))

    if caption:
        caption_paragraph = document.add_paragraph()
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_run = caption_paragraph.add_run(caption)
        set_run_font(caption_run, "SimSun", size=9)


def add_block(document: Document, block: str, body_font: str, title_font: str) -> None:
    heading_match = re.match(r"^(#{1,3})\s+(.*)$", block)
    if heading_match:
        level = len(heading_match.group(1))
        text = strip_inline_markdown(heading_match.group(2))
        paragraph = document.add_paragraph()
        run = paragraph.add_run(text)
        set_run_font(run, title_font, size={1: 18, 2: 16, 3: 14}[level], bold=True)
        return

    bullet_lines = [line.strip() for line in block.splitlines() if line.strip()]
    if bullet_lines and all(line.startswith(("- ", "* ")) for line in bullet_lines):
        for line in bullet_lines:
            paragraph = document.add_paragraph(style="List Bullet")
            run = paragraph.add_run(strip_inline_markdown(line[2:].strip()))
            set_run_font(run, body_font, size=12)
        return

    for line in [line.strip() for line in block.splitlines() if line.strip()]:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(strip_inline_markdown(line))
        set_run_font(run, body_font, size=12)


def build_document(payload: dict, body_font: str, title_font: str) -> Document:
    content = str(payload.get("content", "")).strip()
    if not content:
        raise SystemExit("Payload content is required.")

    blocks = split_blocks(content)
    grouped_images = normalize_images(extract_images(payload), len(blocks))

    document = Document()
    title = str(payload.get("title", "")).strip()
    if title:
        title_paragraph = document.add_paragraph()
        title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_paragraph.add_run(title)
        set_run_font(title_run, title_font, size=20, bold=True)

    for image in grouped_images["cover"]:
        add_image(document, image["path"], image["caption"])

    for index, block in enumerate(blocks, start=1):
        add_block(document, block, body_font, title_font)
        for image in grouped_images["after_paragraph"].get(index, []):
            add_image(document, image["path"], image["caption"], width_inches=5.0)

    for image in grouped_images["end"]:
        add_image(document, image["path"], image["caption"], width_inches=5.0)

    return document


def resolve_output_path(input_path: Path, explicit_output: str | None) -> Path:
    if explicit_output:
        return Path(explicit_output)
    return input_path.with_suffix(".docx")


def main() -> None:
    args = parse_args()
    input_path = Path(args.input)
    payload = load_payload(input_path)
    document = build_document(payload, args.body_font, args.title_font)
    output_path = resolve_output_path(input_path, args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    print(output_path)


if __name__ == "__main__":
    main()
